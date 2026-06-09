"""Integration tests: plan building, rendering, and diffing."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from obelisk_api.domain.models import Athlete, TrainingBlock, TrainingDay
from obelisk_api.services.diff import compute_plan_diff, format_diff
from obelisk_api.services.planbuilder import (
    build_default_plan,
    build_endurance_block_plan,
    build_linear_novice_plan,
)
from obelisk_api.services.render import render_plan

pytestmark = pytest.mark.integration

PROJECT_ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def josh() -> Athlete:
    data = json.loads((PROJECT_ROOT / "data" / "sample_athlete.json").read_text(encoding="utf-8"))
    return Athlete.model_validate(data)


@pytest.fixture
def plan(josh: Athlete):  # type: ignore[no-untyped-def]
    return build_default_plan(
        josh,
        title="Cycle 1 — Aerobic Base + Submaximal Strength",
        goals=["1. Bench priority.", "2. Pull-ups.", "3. Finger strength."],
        priority_lifts=["bench_press"],
    )


def test_plan_training_maxes_match_reference(plan) -> None:  # type: ignore[no-untyped-def]
    tms = {row.lift_key: row for row in plan.training_maxes}
    assert tms["back_squat"].tm_wave1 == 220
    assert tms["back_squat"].tm_wave2 == 230
    assert tms["back_squat"].tm_wave3 == 240
    assert tms["bench_press"].tm_wave1 == 155
    assert tms["bench_press"].tm_wave2 == 160
    assert tms["bench_press"].tm_wave3 == 165
    assert tms["deadlift"].tm_wave1 == 285
    assert tms["strict_press"].tm_wave1 == 105
    assert tms["front_squat"].tm_wave1 == 210
    assert tms["power_clean"].tm_wave1 == 155


def test_plan_has_three_waves_and_maf_cap(plan) -> None:  # type: ignore[no-untyped-def]
    assert len(plan.waves) == 3
    assert plan.maf_cap_bpm == 146
    assert plan.waves[1].training_maxes["back_squat"] == 230


def test_bench_priority_adds_pressing_volume(plan) -> None:  # type: ignore[no-untyped-def]
    monday = plan.day_template[0]
    labels = " ".join(b.label.lower() for b in monday.blocks)
    assert "bench priority" in labels


def test_render_produces_expected_sheets_and_weights(plan, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    out = tmp_path / "cycle_plan.xlsx"
    result = render_plan(plan, str(out))
    assert Path(result).exists()

    wb = load_workbook(result)
    assert wb.sheetnames[0] == "Overview"
    assert "Wave 1 (Wk 1-4)" in wb.sheetnames
    assert "Wave 3 (Wk 9-12)" in wb.sheetnames

    wave1 = wb["Wave 1 (Wk 1-4)"]
    col_e = {cell.value for row in wave1.iter_rows() for cell in row if cell.column == 5}
    # Squat wave1 week1: 65/75/85% of 220 -> 145/165/185.
    assert {145, 165, 185}.issubset(col_e)
    # Deadlift wave1 week1 top set: 85% of 285 = 242 -> 240.
    assert 240 in col_e


def test_diff_detects_friday_trap_bar_swap(plan) -> None:  # type: ignore[no-untyped-def]
    before = plan
    # Swap Friday deadlift -> trap bar for wave 1 only via a days_override.
    new_days: list[TrainingDay] = [d.model_copy(deep=True) for d in plan.day_template]
    friday = new_days[4]
    friday.blocks[0] = TrainingBlock(
        label="Trap Bar Deadlift — main", kind="main", lift_key="trap_bar_deadlift"
    )
    after = plan.model_copy(deep=True)
    after.waves[0].days_override = new_days
    after.waves[0].training_maxes["trap_bar_deadlift"] = 285

    diff = compute_plan_diff(before, after)
    assert not diff.is_empty()
    paths = " ".join(c.path for c in diff.changes)
    assert "waves[0]" in paths
    assert "trap" in format_diff(diff).lower()


def test_diff_empty_when_identical(plan) -> None:  # type: ignore[no-untyped-def]
    diff = compute_plan_diff(plan, plan.model_copy(deep=True))
    assert diff.is_empty()
    assert "No changes" in format_diff(diff)


def _athlete(name: str) -> Athlete:
    data = json.loads((PROJECT_ROOT / "data" / f"sample_{name}.json").read_text(encoding="utf-8"))
    return Athlete.model_validate(data)


def test_novice_plan_is_differentiated(tmp_path: Path) -> None:
    sarah = _athlete("sarah")
    assert sarah.estimated_1rm is None  # loads despite null 1RMs
    p = build_linear_novice_plan(sarah, "Sarah TM", ["finish OCR", "lose 10 lb"])
    assert p.program_model == "linear_novice"
    assert p.training_maxes == []  # no fabricated TMs
    assert p.hangboard_notes == []  # not a climber
    assert p.maf_cap_bpm == 148  # 180 - 32
    assert len(p.waves) == 4 and all(w.phase for w in p.waves)
    assert p.nutrition.protein_g == 132  # 0.8 g/lb, not Josh's 1.0
    # Renders without the strength-only sections.
    out = render_plan(p, str(tmp_path / "sarah.xlsx"))
    assert load_workbook(out).sheetnames[0] == "Overview"


def test_marathon_plan_is_differentiated(tmp_path: Path) -> None:
    marcus = _athlete("marcus")
    p = build_endurance_block_plan(
        marcus, "Marcus TM", ["sub-3"], peak_mileage=70, long_run_peak_mi=22
    )
    assert p.program_model == "marathon_block"
    assert p.training_maxes == [] and p.hangboard_notes == []
    assert p.maf_cap_bpm == 152  # 180 - 28
    assert (p.waves[0].phase or "").startswith("Base")
    # Endurance carb-loading: far above the hybrid 2.0 g/lb cap.
    assert p.nutrition.carbs_by_day["hard"] >= 4 * marcus.bodyweight_lb
    assert p.nutrition.protein_g == 109  # 0.75 g/lb
    render_plan(p, str(tmp_path / "marcus.xlsx"))


def test_render_cycle_plan_xlsx_tool_accepts_dict(plan, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    # The tool wrapper validates a plain dict via the Pydantic model, then renders.
    from obelisk_api.tools.deterministic import render_cycle_plan_xlsx

    out = tmp_path / "from_dict.xlsx"
    result = render_cycle_plan_xlsx(plan.model_dump(), str(out))
    assert Path(result).exists()
    assert load_workbook(result).sheetnames[0] == "Overview"

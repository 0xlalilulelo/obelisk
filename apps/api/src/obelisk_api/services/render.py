"""Render a :class:`~obelisk_api.domain.models.CyclePlan` to an Excel workbook.

Adapted from the hand-written ``build_cycle1.py`` reference: same visual layout
(Overview tab + one tab per wave, colour legend, computed-vs-input cells) but
fully driven by the plan object instead of hard-coded data. Every prescribed
weight is computed here from the wave training max and the loading percentage,
so the numbers are deterministic and never invented by the LLM.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from obelisk_api.domain.models import CyclePlan, TrainingBlock, Wave, WorkSet
from obelisk_api.tools.deterministic import round_to_5

ACCENT = "1F4E79"
HEADER_FILL = "D9E2F3"
INPUT_FILL = "FFF2CC"
COMPUTED_FILL = "E2EFDA"
DELOAD_FILL = "FCE4D6"
DAY_FILL = "F2F2F2"
SUBTLE = "595959"

_thin = Side(border_style="thin", color="BFBFBF")
BORDER = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color=ACCENT)
H2_FONT = Font(name="Calibri", size=14, bold=True, color=ACCENT)
DAY_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
LABEL_FONT = Font(name="Calibri", size=11, bold=True)
TEXT_FONT = Font(name="Calibri", size=11)
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color=SUBTLE)
SMALL_FONT = Font(name="Calibri", size=10)
BOLD_COMPUTED_FONT = Font(name="Calibri", size=11, bold=True)

INPUT = PatternFill("solid", fgColor=INPUT_FILL)
COMPUTED = PatternFill("solid", fgColor=COMPUTED_FILL)
HEADER = PatternFill("solid", fgColor=ACCENT)
SUB_HEADER = PatternFill("solid", fgColor=HEADER_FILL)
DELOAD = PatternFill("solid", fgColor=DELOAD_FILL)
DAY = PatternFill("solid", fgColor=DAY_FILL)


def _reps_text(s: WorkSet) -> str:
    return f"{s.reps}+" if s.amrap else str(s.reps)


def _merge_text(
    ws: Worksheet, row: int, col_start: int, col_end: int, value: str, font: Font
) -> None:
    cell = ws.cell(row=row, column=col_start, value=value)
    cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=col_start, end_row=row, end_column=col_end)


def _build_overview(ws: Worksheet, plan: CyclePlan) -> None:
    for col, width in [(1, 4), (2, 22), (3, 14), (4, 14), (5, 14), (6, 14), (7, 14), (8, 40)]:
        ws.column_dimensions[get_column_letter(col)].width = width

    _merge_text(ws, 2, 2, 8, plan.title, TITLE_FONT)
    _merge_text(ws, 3, 2, 8, plan.subtitle, NOTE_FONT)

    r = 5
    _merge_text(ws, r, 2, 8, "Cycle goals (in priority order)", H2_FONT)
    r += 1
    for goal in plan.goals:
        _merge_text(ws, r, 2, 8, goal, TEXT_FONT)
        r += 1
    r += 1

    # Training maxes (strength plans only)
    if plan.training_maxes:
        _merge_text(ws, r, 2, 8, "Training Maxes (TM = 85% of est 1RM)", H2_FONT)
        r += 1
        headers = ["Lift", "Est. 1RM", "TM Wave 1", "TM Wave 2", "TM Wave 3", "Increment", "Notes"]
        for i, h in enumerate(headers):
            c = ws.cell(row=r, column=2 + i, value=h)
            c.font = LABEL_FONT
            c.fill = SUB_HEADER
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        r += 1
        for row in plan.training_maxes:
            values: list[str | int] = [
                row.display_name,
                row.est_1rm,
                row.tm_wave1,
                row.tm_wave2,
                row.tm_wave3,
                f"+{row.increment}/wave",
                row.notes,
            ]
            for i, v in enumerate(values):
                c = ws.cell(row=r, column=2 + i, value=v)
                c.font = LABEL_FONT if i == 0 else TEXT_FONT
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        r += 1

    # Weekly schedule
    if plan.weekly_schedule:
        _merge_text(ws, r, 2, 8, "Weekly Schedule", H2_FONT)
        r += 1
        for sched in plan.weekly_schedule:
            day_cell = ws.cell(row=r, column=2, value=sched.cells[0])
            day_cell.font = LABEL_FONT
            day_cell.fill = DAY
            day_cell.border = BORDER
            day_cell.alignment = Alignment(horizontal="center")
            _merge_text(ws, r, 3, 8, sched.cells[1], TEXT_FONT)
            r += 1
        r += 1

    # Loading scheme (only when main lifts use %-of-TM loading)
    if any(w.main_sets for w in plan.loading_scheme):
        _merge_text(ws, r, 2, 8, "Loading scheme (per 4-week wave)", H2_FONT)
        r += 1
        sched_headers = ["Week", "Set 1", "Set 2", "Set 3 (top)", "Notes"]
        for i, h in enumerate(sched_headers):
            c = ws.cell(row=r, column=2 + i, value=h)
            c.font = LABEL_FONT
            c.fill = SUB_HEADER
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
        r += 1
        for week in plan.loading_scheme:
            label_cell = ws.cell(row=r, column=2, value=week.label)
            label_cell.font = LABEL_FONT
            label_cell.border = BORDER
            if week.is_deload:
                label_cell.fill = DELOAD
            for i, s in enumerate(week.main_sets[:3]):
                c = ws.cell(row=r, column=3 + i, value=f"{int(s.pct * 100)}% × {_reps_text(s)}")
                c.font = TEXT_FONT
                c.border = BORDER
                c.alignment = Alignment(horizontal="center")
            note = "NO AMRAP — full rest" if week.is_deload else "Top set AMRAP — leave 1-2 RIR"
            note_cell = ws.cell(row=r, column=6, value=note)
            note_cell.font = SMALL_FONT
            note_cell.border = BORDER
            r += 1
        r += 2

    # Aerobic / endurance progression
    if plan.aerobic_build:
        _merge_text(
            ws, r, 2, 8, f"Endurance progression (MAF cap = {plan.maf_cap_bpm} bpm)", H2_FONT
        )
        r += 1
        for i, h in enumerate(plan.aerobic_headers[:5]):
            c = ws.cell(row=r, column=2 + i, value=h)
            c.font = LABEL_FONT
            c.fill = SUB_HEADER
            c.border = BORDER
            c.alignment = Alignment(horizontal="center")
        r += 1
        for aero in plan.aerobic_build:
            cells = aero.cells
            is_deload = "deload" in cells[0].lower() or "taper" in cells[0].lower()
            label_cell = ws.cell(row=r, column=2, value=cells[0])
            label_cell.font = LABEL_FONT
            label_cell.border = BORDER
            if is_deload:
                label_cell.fill = DELOAD
            for i, v in enumerate(cells[1:5]):
                c = ws.cell(row=r, column=3 + i, value=v)
                c.font = SMALL_FONT if i == 3 else TEXT_FONT
                c.border = BORDER
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            r += 1
        r += 2

    # Hangboard (climbers only)
    if plan.hangboard_notes:
        _merge_text(ws, r, 2, 8, "Hangboard schedule (Hörst 7-53 protocol)", H2_FONT)
        r += 1
        for note in plan.hangboard_notes:
            _merge_text(ws, r, 2, 8, "• " + note, SMALL_FONT)
            r += 1
        r += 1

    # Nutrition
    _merge_text(ws, r, 2, 8, "Nutrition profile for this cycle", H2_FONT)
    r += 1
    for note in plan.nutrition.notes:
        _merge_text(ws, r, 2, 8, "• " + note, SMALL_FONT)
        r += 1
    r += 1

    # Legend
    _merge_text(ws, r, 2, 8, "Legend", H2_FONT)
    r += 1
    legend = [
        (INPUT, "Yellow", "Fill in your actual loads / times / RPE"),
        (COMPUTED, "Green", "Auto-computed prescribed load (rounded to 5 lb)"),
        (DELOAD, "Orange", "Deload week — no AMRAP, easy aerobic, full recovery"),
    ]
    for fill, label, desc in legend:
        c = ws.cell(row=r, column=2, value=label)
        c.fill = fill
        c.font = LABEL_FONT
        _merge_text(ws, r, 3, 8, desc, TEXT_FONT)
        r += 1


def _render_strength_block(
    ws: Worksheet, r: int, block: TrainingBlock, sets: list[WorkSet], tm: int
) -> int:
    """Render a main/accessory block (label + per-set table). Returns next row."""
    ws.cell(row=r, column=2, value=block.label).font = TEXT_FONT
    ws.cell(row=r, column=2).border = BORDER
    sub = ["% TM", "Reps", "Prescribed lb", "Actual lb", "Actual reps", "RPE / notes"]
    for i, h in enumerate(sub):
        c = ws.cell(row=r, column=3 + i, value=h)
        c.font = SMALL_FONT
        c.fill = SUB_HEADER
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    r += 1
    for set_i, s in enumerate(sets, start=1):
        ws.cell(row=r, column=2, value=f"  Set {set_i}").font = SMALL_FONT
        ws.cell(row=r, column=2).border = BORDER
        pct_cell = ws.cell(row=r, column=3, value=f"{int(s.pct * 100)}%")
        pct_cell.font = SMALL_FONT
        pct_cell.border = BORDER
        pct_cell.alignment = Alignment(horizontal="center")
        reps_cell = ws.cell(row=r, column=4, value=_reps_text(s))
        reps_cell.font = SMALL_FONT
        reps_cell.border = BORDER
        reps_cell.alignment = Alignment(horizontal="center")
        presc = ws.cell(row=r, column=5, value=round_to_5(tm * s.pct))
        presc.font = BOLD_COMPUTED_FONT
        presc.fill = COMPUTED
        presc.border = BORDER
        presc.alignment = Alignment(horizontal="center")
        for col in (6, 7, 8):
            ic = ws.cell(row=r, column=col)
            ic.fill = INPUT
            ic.border = BORDER
        r += 1
    return r


def _render_note_block(ws: Worksheet, r: int, block: TrainingBlock) -> int:
    ws.cell(row=r, column=2, value=block.label).font = TEXT_FONT
    ws.cell(row=r, column=2).border = BORDER
    note_cell = ws.cell(row=r, column=3, value=block.note or "")
    note_cell.font = SMALL_FONT
    note_cell.border = BORDER
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
    ic = ws.cell(row=r, column=8)
    ic.fill = INPUT
    ic.border = BORDER
    return r + 1


def _build_wave_tab(ws: Worksheet, plan: CyclePlan, wave: Wave) -> None:
    for col, width in [(1, 3), (2, 22), (3, 11), (4, 9), (5, 13), (6, 9), (7, 9), (8, 28)]:
        ws.column_dimensions[get_column_letter(col)].width = width

    first_week = (wave.wave_num - 1) * 4 + 1
    last_week = wave.wave_num * 4
    _merge_text(
        ws, 2, 2, 8, f"Wave {wave.wave_num} — Weeks {first_week} through {last_week}", TITLE_FONT
    )
    tm = wave.training_maxes
    short = {
        "back_squat": "Squat",
        "bench_press": "Bench",
        "deadlift": "DL",
        "strict_press": "Press",
        "front_squat": "FS",
        "power_clean": "PC",
        "trap_bar_deadlift": "Trap DL",
    }
    tm_line = " • ".join(
        f"{short.get(row.lift_key, row.display_name)} {tm.get(row.lift_key, '-')}"
        for row in plan.training_maxes
    )
    if tm_line:
        _merge_text(ws, 3, 2, 8, f"TMs this wave: {tm_line}", NOTE_FONT)
    elif wave.phase:
        _merge_text(ws, 3, 2, 8, f"Phase: {wave.phase}", NOTE_FONT)

    r = 5
    for week_in_wave in range(1, 5):
        global_week = first_week + week_in_wave - 1
        # The per-week loading scheme drives main-lift sets when present; week
        # labels prefer the wave's own labels (endurance/novice phases differ
        # across waves) and fall back to the shared loading scheme.
        scheme = (
            plan.loading_scheme[week_in_wave - 1]
            if week_in_wave - 1 < len(plan.loading_scheme)
            else None
        )
        if wave.week_labels and week_in_wave - 1 < len(wave.week_labels):
            label = wave.week_labels[week_in_wave - 1]
        elif scheme is not None:
            label = scheme.label
        else:
            label = f"Week {global_week}"
        is_deload = (
            (scheme.is_deload if scheme else False)
            or "deload" in label.lower()
            or "taper" in label.lower()
        )
        # Week header
        wc = ws.cell(row=r, column=2, value=f"Week {global_week} — {label}")
        wc.font = DAY_FONT
        wc.fill = DELOAD if is_deload else HEADER
        wc.alignment = Alignment(horizontal="center")
        wc.border = BORDER
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        r += 1

        for day in plan.days_for_week(global_week):
            dc = ws.cell(row=r, column=2, value=f"{day.day} — {day.session}")
            dc.font = LABEL_FONT
            dc.fill = DAY
            dc.border = BORDER
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            r += 1

            for block in day.blocks:
                if (
                    block.kind == "main"
                    and block.lift_key is not None
                    and scheme
                    and scheme.main_sets
                ):
                    block_tm = wave.training_maxes.get(block.lift_key, 0)
                    r = _render_strength_block(ws, r, block, scheme.main_sets, block_tm)
                elif (
                    block.kind == "accessory"
                    and block.lift_key is not None
                    and block.accessory_sets
                ):
                    block_tm = wave.training_maxes.get(block.lift_key, 0)
                    r = _render_strength_block(ws, r, block, block.accessory_sets, block_tm)
                else:
                    r = _render_note_block(ws, r, block)
            r += 1  # space between days
        r += 1  # space between weeks


def render_plan(plan: CyclePlan, output_path: str) -> str:
    """Render the plan to ``output_path`` and return the path."""
    wb = Workbook()
    overview = wb.active
    assert overview is not None
    overview.title = "Overview"
    _build_overview(overview, plan)

    for wave in plan.waves:
        ws = wb.create_sheet(
            f"Wave {wave.wave_num} (Wk {(wave.wave_num - 1) * 4 + 1}-{wave.wave_num * 4})"
        )
        _build_wave_tab(ws, plan, wave)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out)
